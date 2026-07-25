import sys
sys.path.insert(0, '/Users/shubham_kumar/Downloads/Project_1_UG_01')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation('/Users/shubham_kumar/Downloads/Project_1_UG_01/Review 1 - 22 07 2026.pptx')

print('=== SLIDE DIMENSIONS ===')
print(f'Width: {prs.slide_width}, Height: {prs.slide_height}')
print(f'Width (inches): {prs.slide_width / 914400}, Height (inches): {prs.slide_height / 914400}')

print('\n=== SLIDE LAYOUTS ===')
for i, layout in enumerate(prs.slide_layouts):
    print(f'  Layout {i}: {layout.name}')
    for ph in layout.placeholders:
        print(f'    Placeholder {ph.placeholder_format.idx}: {ph.name}, type={ph.placeholder_format.type}')

print('\n=== SLIDES ===')
for si, slide in enumerate(prs.slides):
    print(f'\n--- Slide {si+1} ---')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        print(f'  Shape: type={shape.shape_type}, name="{shape.name}", pos=({shape.left},{shape.top}), size=({shape.width},{shape.height})')
        if hasattr(shape, 'text') and shape.text:
            text_preview = shape.text[:200].replace(chr(10), ' | ')
            print(f'    Text: "{text_preview}"')
        if shape.has_text_frame:
            for pi, para in enumerate(shape.text_frame.paragraphs):
                alignment = para.alignment
                print(f'    Para {pi}: alignment={alignment}, level={para.level}')
                for ri, run in enumerate(para.runs):
                    font = run.font
                    color_val = None
                    try:
                        color_val = str(font.color.rgb) if font.color and font.color.rgb else None
                    except:
                        pass
                    print(f'      Run {ri}: text="{run.text[:80]}", font={font.name}, size={font.size}, bold={font.bold}, ital={font.italic}, color={color_val}')
