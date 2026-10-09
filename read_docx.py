import docx
import sys

def read_docx(file_path):
    doc = docx.Document(file_path)
    for para in doc.paragraphs:
        print(para.text)
    
    for table in doc.tables:
        print("\n--- TABLE ---")
        for row in table.rows:
            print(" | ".join([cell.text.replace('\n', ' ') for cell in row.cells]))
        print("-------------\n")

if __name__ == '__main__':
    read_docx(sys.argv[1])
